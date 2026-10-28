"""Convert the Happier Lives Institute's wellbeing models of the Against Malaria Foundation and
GiveDirectly into charityevals models. They are second models for charities GiveWell also models.

AMF ("The elephant in the bednet", 2022, sheet 1RrBu...): WELLBYs per $1,000 = deaths averted per
$1,000 (from GiveWell's AMF model) × wellbeing value of each life saved, plus wellbeing gains from
better health and income. The value of a life saved depends on a philosophical choice HLI leaves
open: the "neutral point" on the 0–10 life-satisfaction scale above which life counts as good.
HLI publishes results for neutral points 0–5 under deprivationism; the model takes the neutral
point as an input over that range (best guess 2, our pick) and reproduces HLI's row for it.

GiveDirectly: HLI's current estimate, 7.55 WELLBYs per $1,000, from its 2024 summary sheet.

    python3 scripts/convert-hli-amf.py amf.xlsx smfb.xlsx --out models/global-health
"""
import argparse, json
from pathlib import Path
import openpyxl

p = argparse.ArgumentParser()
p.add_argument("amf_xlsx"); p.add_argument("summary_xlsx"); p.add_argument("--out", required=True)
a = p.parse_args()
V = openpyxl.load_workbook(a.amf_xlsx, data_only=True)
S = openpyxl.load_workbook(a.summary_xlsx, data_only=True).worksheets[0]
cea, inp, gw = V["AMF SWB CEA"], V["Inputs to AMF CEA"], V["GiveWell Numbers"]
ROW = 16    # deprivationism, neutral point 2, with HLI's (upper-bound) improving effects
SHEET = "https://docs.google.com/spreadsheets/d/1RrBuiPVgL-t8hlr6EqkqABiaqdHMGkpvfeiqiiX49LU/edit"
REPORT = "https://forum.effectivealtruism.org/posts/JgqEqsa6iAtqGLYmw/the-elephant-in-the-bednet-the-importance-of-philosophy-when-1"
SUMMARY = "https://docs.google.com/spreadsheets/d/1nwTW81QYNTSNZDKFj8ozwECARfILEktJp8uQGwSPBeU/edit"
src = lambda cell, title="HLI, AMF wellbeing CEA": {"title": title, "url": SHEET, "cell": cell}
v = lambda ws, c: ws[c].value
assert v(cea, f"P{ROW}") == 2 and v(cea, f"D{ROW}") == "deprivationist"
np_hi = v(cea, "P19"); assert np_hi == 5

amf_rows = [
    {"id": "deaths_averted_per_1000", "section": "saving", "label": "Deaths averted per $1,000 (from GiveWell's AMF model)", "format": "number",
     "value": v(gw, "B8"), "kind": "estimate", "source": src("GiveWell Numbers!B8"),
     "explainer": "Under-five and over-five deaths averted per $1,000, which HLI takes from its copy of GiveWell's 2022 AMF model. "
                  "GiveWell's current model (the other model of this charity) gives a different figure; HLI's analysis predates it."},
    {"id": "avg_age_at_death", "section": "saving", "label": "Average age of those saved", "unit": "years", "format": "number",
     "value": v(inp, "C13"), "kind": "estimate", "source": src("Inputs to AMF CEA!C13"),
     "explainer": "Deaths averted weighted between under-fives (average age 1.5) and over-fives (46.1)."},
    {"id": "life_expectancy", "section": "saving", "label": "Life expectancy used", "unit": "years", "format": "number",
     "value": v(cea, f"U{ROW}"), "kind": "estimate", "source": src("GiveWell Numbers!B12"), "explainer": "Average life expectancy in AMF's countries, from GiveWell's model."},
    {"id": "years_saved", "section": "saving", "label": "Years of life saved per death averted", "format": "number",
     "formula": "life_expectancy - avg_age_at_death", "expected": v(cea, f"T{ROW}")},
    {"id": "avg_life_satisfaction", "section": "saving", "label": "Average life satisfaction in AMF's countries (0–10)", "format": "number",
     "value": v(cea, f"S{ROW}"), "kind": "observed", "source": src("GiveWell Numbers!B11"), "explainer": "From the Gallup World Poll, as used by HLI."},
    {"id": "neutral_point", "section": "saving", "label": "Neutral point on the 0–10 life-satisfaction scale", "format": "number",
     "value": {"low": 0, "best": 2, "high": np_hi}, "kind": "judgment", "source": src(f"AMF SWB CEA!P14:P19"),
     "explainer": "The level of life satisfaction at which life is neither good nor bad for the person living it. Below it, extra years "
                  "count as bad; above it, good. HLI shows results for 0 to 5 and picks none; the best guess of 2 is ours, close to the "
                  "value some survey respondents report. This single choice moves AMF's result more than any other input."},
    {"id": "net_wellbeing", "section": "saving", "label": "Life satisfaction above the neutral point", "format": "number",
     "formula": "avg_life_satisfaction - neutral_point", "expected": v(cea, f"R{ROW}")},
    {"id": "wellbys_from_years", "section": "saving", "label": "WELLBYs from the years of life saved", "format": "number",
     "formula": "years_saved * net_wellbeing", "expected": v(cea, f"O{ROW}")},
    {"id": "ls_growth", "section": "saving", "label": "Yearly rise in life satisfaction over time", "format": "number",
     "value": v(inp, "C19"), "kind": "estimate", "source": src("Inputs to AMF CEA!C19"),
     "explainer": "Life satisfaction in these countries has been rising; later years of a life saved are expected to be better."},
    {"id": "wellbys_from_growth", "section": "saving", "label": "WELLBYs from rising life satisfaction", "format": "number",
     "formula": "ls_growth * years_saved * years_saved * 0.5", "expected": v(cea, f"N{ROW}")},
    {"id": "grief_effect", "section": "saving", "label": "Effect of a child's death on a family member's life satisfaction", "format": "number",
     "value": v(inp, "C3"), "kind": "estimate", "source": src("Inputs to AMF CEA!C3"), "explainer": "From Oswald and Powdthavee, via HLI."},
    {"id": "grief_years", "section": "saving", "label": "Years the grief lasts", "format": "number",
     "value": v(inp, "C4"), "kind": "estimate", "source": src("Inputs to AMF CEA!C4"), "explainer": "HLI's estimate of how long the effect of bereavement lasts."},
    {"id": "grieving_household", "section": "saving", "label": "Other household members", "format": "number",
     "value": v(inp, "C5"), "kind": "observed", "source": src("Inputs to AMF CEA!C5"), "explainer": "Average household size in AMF's countries minus one, from GiveWell's model."},
    {"id": "wellbys_from_grief", "section": "saving", "label": "WELLBYs from grief averted", "format": "number",
     "formula": "abs(0.5 * grief_effect * grief_years) * grieving_household", "expected": v(cea, f"M{ROW}")},
    {"id": "wellbys_per_life", "section": "saving", "label": "WELLBYs per life saved", "format": "number",
     "formula": "wellbys_from_grief + wellbys_from_years + wellbys_from_growth", "expected": v(cea, f"L{ROW}")},
    {"id": "wellbys_saving_per_1000", "section": "saving", "label": "WELLBYs from saving lives per $1,000", "format": "number",
     "formula": "deaths_averted_per_1000 * wellbys_per_life", "expected": v(cea, f"I{ROW}")},
    {"id": "wellbys_improving_per_1000", "section": "improving", "label": "WELLBYs from better health and income per $1,000", "format": "number",
     "value": {"low": 0, "best": v(cea, f"J{ROW}"), "high": v(cea, f"J{ROW}")}, "kind": "estimate", "source": src("Improving (Upper B)!B103"),
     "explainer": "Less malaria illness and higher later incomes, including community spillovers at 75% net coverage. This is HLI's upper bound; "
                  "its lower-bound tab is broken in the published sheet (#REF!), so the range runs down to HLI's 'just life saving' version, which leaves it out."},
    {"id": "wellbys_per_1000", "section": "bottom", "label": "WELLBYs per $1,000", "format": "number",
     "formula": "wellbys_saving_per_1000 + wellbys_improving_per_1000", "expected": v(cea, f"H{ROW}")},
    {"id": "wellbys_per_dollar", "section": "bottom", "label": "WELLBYs per dollar", "format": "small", "formula": "wellbys_per_1000 / 1000"},
]
amf = {
    "id": "against-malaria-foundation-hli", "name": "AMF bed nets, wellbeing analysis", "kind": "charity", "cause": "global-health",
    "charity": "against-malaria-foundation", "owner": "hli",
    "summary": "The Happier Lives Institute's analysis of AMF in wellbeing terms (WELLBYs). Its result depends heavily on a philosophical choice "
               "HLI leaves open: the neutral point above which a life counts as good. HLI found AMF less cost-effective than StrongMinds under "
               "almost all views it considered.",
    "source": {"publisher": "Happier Lives Institute", "title": "The elephant in the bednet: the importance of philosophy when choosing between extending and improving lives",
               "url": REPORT, "sheet": SHEET, "published": "2022-11-01", "retrieved": "2026-10-27", "license": "Unclear: check HLI's terms before republishing"},
    "evaluator": {"name": "HLI", "assumptions": {"count_salary_sacrifice": 0, "full_credit": 0}},
    "sections": [{"id": "saving", "title": "Value of the lives saved", "explainer": "Deaths averted, each worth the wellbeing of the years saved, plus grief averted for the family."},
                 {"id": "improving", "title": "Better health and income"}, {"id": "bottom", "title": "Bottom line"}],
    "rows": amf_rows,
    "output": {"row": "wellbys_per_dollar", "unit": "wellby", "label": "WELLBYs per dollar"},
    "resources": [
        {"title": "HLI: The elephant in the bednet (2022)", "url": REPORT, "type": "evaluation"},
        {"title": "HLI's AMF wellbeing spreadsheet", "url": SHEET, "type": "model", "note": "The model this page reproduces."},
        {"title": "HLI's companion AMF spreadsheet", "url": "https://docs.google.com/spreadsheets/d/1NMAU-a1X4vqjodjI6kf8KnUyCJaK9uyNvXWj5VetDZw/edit", "type": "model"},
    ],
}
gd = {
    "id": "givedirectly-hli", "name": "GiveDirectly cash transfers, wellbeing analysis", "kind": "charity", "cause": "global-health",
    "charity": "givedirectly", "owner": "hli",
    "summary": "The Happier Lives Institute's estimate of GiveDirectly's effect on wellbeing, from its meta-analysis of cash transfers' effects on "
               "life satisfaction and mental health, including spillovers to the household.",
    "source": {"publisher": "Happier Lives Institute", "title": "Cash transfers cost-effectiveness analysis",
               "url": "https://www.happierlivesinstitute.org/report/cash-transfers-cost-effectiveness-analysis/", "sheet": SUMMARY,
               "published": "2024-11-01", "retrieved": "2026-10-27", "license": "Unclear: check HLI's terms before republishing"},
    "evaluator": {"name": "HLI", "assumptions": {"count_salary_sacrifice": 0, "full_credit": 0}},
    "sections": [{"id": "bottom", "title": "Bottom line"}],
    "rows": [
        {"id": "wellbys_per_1000", "section": "bottom", "label": "WELLBYs per $1,000", "format": "number", "value": S["B45"].value, "kind": "estimate",
         "source": {"title": "HLI summary spreadsheet V4", "url": SUMMARY, "cell": "B45"},
         "explainer": "HLI's current figure for GiveDirectly, which it uses as the benchmark for its other charities. HLI's earlier, more detailed "
                      "model (2022) put it at 8.2 WELLBYs per $1,000. HLI's report explains the meta-analysis behind it; the summary sheet "
                      "publishes only the result, so this model has a single input."},
        {"id": "wellbys_per_dollar", "section": "bottom", "label": "WELLBYs per dollar", "format": "small", "formula": "wellbys_per_1000 / 1000"},
    ],
    "output": {"row": "wellbys_per_dollar", "unit": "wellby", "label": "WELLBYs per dollar"},
    "resources": [
        {"title": "HLI: Cash transfers cost-effectiveness analysis", "url": "https://www.happierlivesinstitute.org/report/cash-transfers-cost-effectiveness-analysis/", "type": "evaluation"},
        {"title": "HLI summary spreadsheet V4", "url": SUMMARY, "type": "model"},
    ],
}
for m in (amf, gd):
    Path(a.out, f"{m['id']}.json").write_text(json.dumps(m, indent=2, ensure_ascii=False) + "\n")
    print(m["id"], len(m["rows"]), "rows")
