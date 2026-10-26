"""Dump a public Google Sheet (or local .xlsx) as JSON lines: one per non-empty row, with each
cell's formula and cached value, so a model import can be written row by row.

    python3 scripts/sheet-dump.py <google-sheet-url-or-xlsx> [--tab NAME] [--cols B:D] > dump.jsonl

Needs openpyxl (pip install openpyxl). Run it in its own scratch folder; it only reads the file.
"""
import argparse, json, re, sys, tempfile, urllib.request
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

p = argparse.ArgumentParser()
p.add_argument("src")
p.add_argument("--tab", action="append", help="tab name (repeatable); default: all tabs")
p.add_argument("--cols", default="A:H", help="column range to dump, e.g. A:F")
p.add_argument("--max-rows", type=int, default=400)
a = p.parse_args()

path = Path(a.src)
m = re.search(r"/spreadsheets/d/([\w-]+)", a.src)
if m:
    path = Path(tempfile.mkdtemp()) / f"{m.group(1)}.xlsx"
    url = f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx"
    urllib.request.urlretrieve(url, path)
    print(f"# downloaded {url}", file=sys.stderr)

formulas = openpyxl.load_workbook(path)
values = openpyxl.load_workbook(path, data_only=True)
c0, c1 = (column_index_from_string(c) for c in a.cols.split(":"))
print(f"# tabs: {formulas.sheetnames}", file=sys.stderr)

for name in a.tab or formulas.sheetnames:
    fws, vws = formulas[name], values[name]
    for r in range(1, min(fws.max_row, a.max_rows) + 1):
        cells = {}
        for c in range(c0, min(c1, fws.max_column) + 1):
            f, v = fws.cell(r, c).value, vws.cell(r, c).value
            if f is None and v is None:
                continue
            f = getattr(f, "text", f)  # array formulas
            cell = {"v": v}
            if isinstance(f, str) and f.startswith("="):
                # Google-only functions arrive wrapped; keep the original text.
                g = re.search(r'__xludf\.DUMMYFUNCTION\("(.*)"\),', f)
                cell["f"] = g.group(1).replace('""', '"') if g else f
            cells[get_column_letter(c)] = cell
        if cells:
            print(json.dumps({"tab": name, "row": r, "cells": cells}, default=str))
